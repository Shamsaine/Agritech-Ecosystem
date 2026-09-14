"""Read-only Excel workbook loading for the pilot import."""

import hashlib
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


@dataclass(frozen=True)
class WorkbookData:
    path: Path
    checksum_sha256: str
    sheets: dict[str, list[dict[str, Any]]]


def calculate_checksum(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_workbook_data(source: str | Path) -> WorkbookData:
    path = Path(source).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Workbook not found: {path}")
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Checkpoint 12 accepts .xlsx workbooks only")

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, list[dict[str, Any]]] = {}
    try:
        for worksheet in workbook.worksheets:
            values = worksheet.iter_rows(values_only=True)
            header_row = next(values, None)
            if header_row is None:
                sheets[worksheet.title] = []
                continue

            headers = [str(value).strip() if value is not None else "" for value in header_row]
            rows: list[dict[str, Any]] = []
            for values_row in values:
                if not any(value not in (None, "") for value in values_row):
                    continue
                rows.append(
                    {
                        header: values_row[index] if index < len(values_row) else None
                        for index, header in enumerate(headers)
                        if header
                    }
                )
            sheets[worksheet.title] = rows
    finally:
        workbook.close()

    return WorkbookData(path, calculate_checksum(path), sheets)